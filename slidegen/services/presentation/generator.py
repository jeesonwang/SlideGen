import asyncio
import json
import uuid
from collections.abc import AsyncGenerator
from functools import partial
from pathlib import Path
from time import perf_counter
from typing import Any

from agno.models.base import Model
from agno.models.message import Message
from loguru import logger
from lxml.etree import Element, fromstring, tostring
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation as PresentationType

from slidegen.schemas.gen_request import ExportFormat, GeneratePresentationRequest
from slidegen.schemas.stream_event import (
    ProgressEvent,
    StepCompletedEvent,
    StepStartedEvent,
    StreamEventT,
    WorkflowCompletedEvent,
    WorkflowErrorEvent,
)
from slidegen.schemas.theme import PresentationTheme, ThemePresets
from slidegen.services.document.markdown import MarkdownDocument
from slidegen.services.presentation.orchestrator import PresentationOrchestrator
from slidegen.services.presentation.pdf_exporter import pdf_exporter
from slidegen.services.slidegen.workflow import get_llm_instance, run_slidegen_workflow, run_slidegen_workflow_stream

AUTO_THEME_PRESETS = {"auto", "auto_theme", "__auto_theme__"}
AUTO_THEME_FALLBACK = "modern_minimalist"


def is_auto_theme_preset(theme_preset: str | None) -> bool:
    return bool(theme_preset and theme_preset.casefold() in AUTO_THEME_PRESETS)


class PresentationGenerator:
    """Generator for creating PowerPoint presentations from user requests"""

    def __init__(self, templates_dir: str | None = None) -> None:
        """Initialize the presentation generator."""
        if templates_dir is None:
            project_root = Path(__file__).parent.parent.parent.parent
            templates_dir = str(project_root / "components" / "templates")

        self.templates_dir = Path(templates_dir)
        self.converter = PresentationOrchestrator()

    def get_template_path(self, template_name: str) -> str:
        """Get the full path for a template by name."""
        template_file = self.templates_dir / f"template_{template_name}.pptx"
        if not template_file.exists():
            raise FileNotFoundError(f"Template '{template_name}' not found at: {template_file}")
        return str(template_file)

    def list_templates(self) -> list[str]:
        """List all available template names."""
        return sorted(file.stem.replace("template_", "") for file in self.templates_dir.glob("template_*.pptx"))

    def _request_theme_context(self, request: GeneratePresentationRequest, generated_content: str | None = None) -> str:
        chunks = [
            request.content,
            request.instructions or "",
            "\n".join(request.slides_markdown or []),
            generated_content or "",
        ]
        return "\n".join(chunk for chunk in chunks if chunk)

    def _available_theme_options(self) -> list[dict[str, str]]:
        return [
            {"id": preset_id, "name": preset.name}
            for preset_id in ThemePresets.list_presets()
            if (preset := ThemePresets.get_preset(preset_id)) is not None
        ]

    def _build_auto_theme_prompt(self, content: str | None) -> str:
        options = self._available_theme_options()
        return (
            "Choose the best PowerPoint visual theme preset for the presentation content.\n"
            'Return JSON only in this shape: {"theme_preset":"preset_id"}.\n'
            "The theme_preset value must be exactly one id from the available presets.\n\n"
            f"Available presets:\n{json.dumps(options, ensure_ascii=False)}\n\n"
            f"Presentation content:\n{(content or '').strip()[:6000]}"
        )

    def _extract_theme_preset_id(self, response_content: Any) -> str | None:
        text = str(response_content or "").strip()
        available = ThemePresets.list_presets()
        if text in set(available):
            return text

        try:
            payload = json.loads(text)
        except json.JSONDecodeError:
            payload = None

        if isinstance(payload, dict):
            candidate = payload.get("theme_preset") or payload.get("preset") or payload.get("id")
            if isinstance(candidate, str) and candidate in available:
                return candidate

        textual_matches = [preset_id for preset_id in available if preset_id in text]
        if len(textual_matches) == 1:
            return textual_matches[0]

        return None

    async def _select_auto_theme_preset(self, content: str | None, llm: Model | Any | None) -> str:
        """Ask the configured LLM to select one existing theme preset id."""
        if llm is None:
            logger.warning("Auto Theme requested without an LLM; using fallback theme preset")
            return AUTO_THEME_FALLBACK

        prompt = self._build_auto_theme_prompt(content)
        message = Message(role="user", content=prompt)

        try:
            response = await llm.aresponse([message])
            response_content = response.content if hasattr(response, "content") else response
            selected_preset = self._extract_theme_preset_id(response_content)
            if selected_preset:
                return selected_preset
            logger.warning(f"Auto Theme LLM returned invalid preset: {response_content!r}")
        except Exception as e:
            logger.warning(f"Auto Theme LLM selection failed: {e}")

        return AUTO_THEME_FALLBACK

    async def _resolve_theme(
        self,
        theme: PresentationTheme | None = None,
        theme_preset: str | None = None,
        auto_content: str | None = None,
        auto_theme_llm: Model | Any | None = None,
    ) -> PresentationTheme | None:
        """
        Resolve theme from either direct theme object or preset name.

        Args:
            theme: Direct theme object (takes priority)
            theme_preset: Theme preset name, or "auto" to infer from content
            auto_content: Topic or slide content used when theme_preset is "auto"
            auto_theme_llm: LLM used to select an existing preset when theme_preset is "auto"

        Returns:
            Resolved theme or None if neither is provided
        """
        if theme:
            return theme
        if theme_preset:
            if is_auto_theme_preset(theme_preset):
                selected_preset = await self._select_auto_theme_preset(auto_content, auto_theme_llm)
                preset_theme = ThemePresets.get_preset(selected_preset)
                if preset_theme:
                    logger.info(f"Auto-selected theme preset: {selected_preset}")
                    return preset_theme

            preset_theme = ThemePresets.get_preset(theme_preset)
            if preset_theme:
                logger.info(f"Using theme preset: {theme_preset}")
                return preset_theme
            else:
                logger.warning(f"Theme preset '{theme_preset}' not found, available: {ThemePresets.list_presets()}")
        return None

    def apply_theme_colors(self, presentation: PresentationType, theme: PresentationTheme) -> None:
        """
        Apply theme colors to a PowerPoint presentation.

        This method modifies the presentation's theme XML directly, ensuring that
        all elements using theme colors are automatically updated.

        Args:
            presentation: The PowerPoint presentation object
            theme: The theme configuration with color mappings

        Example:
            >>> from slidegen.schemas.theme import ThemePresets
            >>> generator = PresentationGenerator()
            >>> prs = Presentation("template.pptx")
            >>> generator.apply_theme_colors(prs, ThemePresets.SUNSET_BOULEVARD)
            >>> prs.save("themed.pptx")
        """
        try:
            # Get the slide master and theme part
            slide_master = presentation.slide_master
            slide_master_part = slide_master.part
            theme_part = slide_master_part.part_related_by(RT.THEME)

            # Parse the theme XML
            theme_xml = fromstring(theme_part.blob)

            # Get the color mappings (only non-None values)
            theme_colors = theme.colors.model_dump_colors()

            # Define the XML namespace
            nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

            # Update each color in the theme
            for color_name, hex_value in theme_colors.items():
                if color_name:
                    try:
                        # Clean hex value first
                        hex_value_clean = hex_value.replace("0x", "").replace("#", "").upper()

                        # Find the color container element (e.g., <a:dk1>, <a:accent1>)
                        color_container = theme_xml.xpath(
                            f"a:themeElements/a:clrScheme/a:{color_name}",
                            namespaces=nsmap,
                        )

                        if not color_container:
                            logger.warning(f"Color container not found for {color_name}")
                            continue

                        color_container = color_container[0]

                        # Try to find srgbClr element first (RGB colors)
                        srgb_elements = color_container.xpath("a:srgbClr", namespaces=nsmap)

                        if srgb_elements:
                            # Update existing srgbClr element
                            srgb_elements[0].set("val", hex_value_clean)
                            logger.debug(f"Updated srgbClr for {color_name} to {hex_value_clean}")
                        else:
                            # Check if it's a sysClr element (system colors like dk1, lt1)
                            sys_elements = color_container.xpath("a:sysClr", namespaces=nsmap)

                            if sys_elements:
                                # Remove the sysClr element
                                color_container.remove(sys_elements[0])
                                logger.debug(f"Removed sysClr for {color_name}")

                            # Create new srgbClr element
                            srgb_element = Element(
                                f"{{{nsmap['a']}}}srgbClr",
                                val=hex_value_clean,
                            )
                            color_container.append(srgb_element)
                            logger.debug(f"Created srgbClr for {color_name} with value {hex_value_clean}")

                    except IndexError:
                        logger.warning(f"Could not find color element for {color_name}")
                    except Exception as e:
                        logger.error(f"Error updating color {color_name}: {e}")

            # Save the modified theme XML back to the theme part
            theme_part._blob = tostring(theme_xml)
            logger.info(f"Successfully applied theme '{theme.name}' with {len(theme_colors)} colors")

        except Exception as e:
            logger.error(f"Failed to apply theme colors: {e}")
            raise

    async def generate_presentation(
        self,
        request: GeneratePresentationRequest,
        output_path: str,
    ) -> str:
        """Generate a PowerPoint presentation from a request."""
        template = self.get_template_path(request.template)
        logger.info(f"Starting presentation generation with template: {template}")

        logger.info("Running slide generation workflow...")
        markdown_doc: MarkdownDocument = await run_slidegen_workflow(request)

        logger.info("Loading template presentation...")
        loop = asyncio.get_event_loop()
        template_prs = await loop.run_in_executor(None, Presentation, template)

        # Apply theme if provided
        auto_theme_llm = await get_llm_instance(request) if is_auto_theme_preset(request.theme_preset) else None
        theme = await self._resolve_theme(
            request.theme,
            request.theme_preset,
            auto_content=self._request_theme_context(request, markdown_doc.text),
            auto_theme_llm=auto_theme_llm,
        )
        if theme:
            logger.info(f"Applying theme: {theme.name}")
            self.apply_theme_colors(template_prs, theme)

        logger.info("Converting Markdown to PowerPoint...")
        presentation = await self.converter.generate(template_prs, markdown_doc, theme=theme)

        # Save PPTX to temp, then conditionally convert
        pptx_tmp = output_path + ".pptx"
        logger.info(f"Saving temporary presentation to: {pptx_tmp}")
        await loop.run_in_executor(None, partial(presentation.save, pptx_tmp))

        if request.export_as == ExportFormat.PDF:
            logger.info(f"Converting PPTX to PDF: {output_path}")
            await loop.run_in_executor(None, partial(pdf_exporter.convert, pptx_tmp, output_path))
            Path(pptx_tmp).unlink(missing_ok=True)
        else:
            logger.info(f"Moving presentation to: {output_path}")
            Path(pptx_tmp).rename(output_path)

        logger.info(f"Successfully generated presentation: {output_path}")
        return output_path

    async def generate_from_markdown(
        self,
        markdown_content: str,
        template_name: str,
        output_path: str,
        export_as: ExportFormat = ExportFormat.PPTX,
        theme: PresentationTheme | None = None,
        theme_preset: str | None = None,
        user_id: uuid.UUID | None = None,
    ) -> str:
        """Generate a presentation directly from markdown content.

        Args:
            markdown_content: The markdown content to convert to presentation
            template_name: Name of the template to use
            output_path: Path to save the generated presentation
            export_as: Export format ("pptx" or "pdf")
            theme: Optional theme to apply to the presentation
            theme_preset: Optional theme preset name
            user_id: User ID for LLM config lookup (required for auto theme)

        Returns:
            The output path of the generated presentation
        """
        started_at = perf_counter()
        template = self.get_template_path(template_name)
        logger.info(
            "Starting presentation generation from markdown: template={}, export_as={}, markdown_chars={}, output_path={}",
            template,
            export_as.value,
            len(markdown_content),
            output_path,
        )

        # Parse markdown content
        parse_started_at = perf_counter()
        markdown_doc = MarkdownDocument(markdown_content)
        logger.info("Markdown parsed in {:.2f}s", perf_counter() - parse_started_at)

        # Load template
        loop = asyncio.get_event_loop()
        load_started_at = perf_counter()
        template_prs = await loop.run_in_executor(None, Presentation, template)
        logger.info("Template loaded in {:.2f}s", perf_counter() - load_started_at)

        # Resolve auto-theme LLM internally
        auto_theme_llm = None
        if is_auto_theme_preset(theme_preset) and user_id is not None:
            logger.info("Resolving LLM instance for auto theme")
            auto_theme_llm = await get_llm_instance(
                GeneratePresentationRequest(
                    content=markdown_content[:6000],
                    template=template_name,
                    export_as=export_as,
                    user_id=user_id,
                )
            )

        # Apply theme colors if provided
        theme_started_at = perf_counter()
        theme = await self._resolve_theme(
            theme,
            theme_preset,
            auto_content=markdown_content,
            auto_theme_llm=auto_theme_llm,
        )
        logger.info("Theme resolved in {:.2f}s", perf_counter() - theme_started_at)
        if theme:
            logger.info(f"Applying theme: {theme.name}")
            self.apply_theme_colors(template_prs, theme)

        # Convert to PPT
        convert_started_at = perf_counter()
        logger.info("Converting Markdown to PowerPoint...")
        presentation = await self.converter.generate(template_prs, markdown_doc, theme=theme)
        logger.info("Markdown converted to PowerPoint in {:.2f}s", perf_counter() - convert_started_at)

        # Save PPTX to temp, then conditionally convert
        pptx_tmp = output_path + ".pptx"
        save_started_at = perf_counter()
        logger.info(f"Saving temporary presentation to: {pptx_tmp}")
        await loop.run_in_executor(None, partial(presentation.save, pptx_tmp))
        logger.info("Temporary presentation saved in {:.2f}s", perf_counter() - save_started_at)

        if export_as == ExportFormat.PDF:
            pdf_started_at = perf_counter()
            logger.info(f"Converting PPTX to PDF: {output_path}")
            await loop.run_in_executor(None, partial(pdf_exporter.convert, pptx_tmp, output_path))
            Path(pptx_tmp).unlink(missing_ok=True)
            logger.info("PDF conversion completed in {:.2f}s", perf_counter() - pdf_started_at)
        else:
            logger.info(f"Moving presentation to: {output_path}")
            Path(pptx_tmp).rename(output_path)

        logger.info(
            "Successfully generated presentation from markdown: {}, total_time={:.2f}s",
            output_path,
            perf_counter() - started_at,
        )
        return output_path

    async def generate_from_markdown_stream(
        self,
        markdown_content: str,
        template_name: str,
        output_path: str,
        export_as: ExportFormat = ExportFormat.PPTX,
        theme: PresentationTheme | None = None,
        theme_preset: str | None = None,
        user_id: uuid.UUID | None = None,
    ) -> AsyncGenerator[StreamEventT, None]:
        """Generate a presentation from markdown with streaming progress events.

        Args:
            markdown_content: The markdown content to convert to presentation
            template_name: Name of the template to use
            output_path: Path to save the generated presentation
            export_as: Export format ("pptx" or "pdf")
            theme: Optional theme to apply to the presentation
            theme_preset: Optional theme preset name
            user_id: User ID for LLM config lookup (required for auto theme)

        Yields:
            Stream events containing conversion progress
        """
        try:
            started_at = perf_counter()
            template = self.get_template_path(template_name)
            logger.info(
                "Starting streaming presentation generation from markdown: template={}, export_as={}, "
                "markdown_chars={}, output_path={}",
                template,
                export_as.value,
                len(markdown_content),
                output_path,
            )

            yield StepStartedEvent(
                step_name="Presentation Export",
                message="Converting markdown to presentation format...",
            )
            logger.info("PPT stream progress 0%: loading template")

            yield ProgressEvent(
                stage="presentation_export",
                progress=0.0,
                message="Loading template...",
            )

            # Load template
            loop = asyncio.get_event_loop()
            load_started_at = perf_counter()
            template_prs = await loop.run_in_executor(None, Presentation, template)
            logger.info("Template loaded in {:.2f}s", perf_counter() - load_started_at)

            # Resolve auto-theme LLM internally
            auto_theme_llm = None
            if is_auto_theme_preset(theme_preset) and user_id is not None:
                logger.info("Resolving LLM instance for auto theme")
                auto_theme_llm = await get_llm_instance(
                    GeneratePresentationRequest(
                        content=markdown_content[:6000],
                        template=template_name,
                        export_as=export_as,
                        user_id=user_id,
                    )
                )

            # Apply theme if provided
            theme = await self._resolve_theme(
                theme,
                theme_preset,
                auto_content=markdown_content,
                auto_theme_llm=auto_theme_llm,
            )
            if theme:
                logger.info("PPT stream progress 10%: applying theme {}", theme.name)
                yield ProgressEvent(
                    stage="presentation_export",
                    progress=10.0,
                    message=f"Applying theme: {theme.name}...",
                )
                self.apply_theme_colors(template_prs, theme)

            logger.info("PPT stream progress 20%: parsing markdown content")
            yield ProgressEvent(
                stage="presentation_export",
                progress=20.0,
                message="Parsing markdown content...",
            )

            # Parse markdown
            parse_started_at = perf_counter()
            markdown_doc = MarkdownDocument(markdown_content)
            logger.info("Markdown parsed in {:.2f}s", perf_counter() - parse_started_at)

            logger.info("PPT stream progress 40%: converting markdown to slides")
            yield ProgressEvent(
                stage="presentation_export",
                progress=40.0,
                message="Converting markdown to slides...",
            )

            # Convert to PPT
            convert_started_at = perf_counter()
            presentation = await self.converter.generate(template_prs, markdown_doc, theme=theme)
            logger.info("Markdown converted to PowerPoint in {:.2f}s", perf_counter() - convert_started_at)

            logger.info("PPT stream progress 80%: saving presentation file")
            yield ProgressEvent(
                stage="presentation_export",
                progress=80.0,
                message="Saving presentation file...",
            )

            # Save PPTX to temp, then conditionally convert
            pptx_tmp = output_path + ".pptx"
            save_started_at = perf_counter()
            await loop.run_in_executor(None, partial(presentation.save, pptx_tmp))
            logger.info("Temporary presentation saved in {:.2f}s", perf_counter() - save_started_at)

            if export_as == ExportFormat.PDF:
                logger.info("PPT stream progress 90%: converting PPTX to PDF")
                yield ProgressEvent(
                    stage="presentation_export",
                    progress=90.0,
                    message="Converting to PDF...",
                )
                await loop.run_in_executor(None, partial(pdf_exporter.convert, pptx_tmp, output_path))
                Path(pptx_tmp).unlink(missing_ok=True)
            else:
                Path(pptx_tmp).rename(output_path)

            logger.info("PPT stream progress 100%: presentation saved successfully")
            yield ProgressEvent(
                stage="presentation_export",
                progress=100.0,
                message="Presentation saved successfully",
            )

            yield StepCompletedEvent(
                step_name="Presentation Export",
                content=output_path,
                message="Presentation file generated successfully",
            )

            logger.info(
                "Successfully generated presentation from markdown: {}, total_time={:.2f}s",
                output_path,
                perf_counter() - started_at,
            )

        except Exception as e:
            logger.exception(f"Failed to generate presentation from markdown: {e}")
            yield WorkflowErrorEvent(
                error=str(e),
                message="Presentation generation from markdown failed",
            )

    async def generate_presentation_stream(
        self,
        request: GeneratePresentationRequest,
        output_path: str,
    ) -> AsyncGenerator[StreamEventT, None]:
        """Generate a PowerPoint presentation with streaming progress events."""
        try:
            template = self.get_template_path(request.template)
            logger.info(f"Starting streaming presentation generation with template: {template}")

            final_content: str | None = None

            async for event in run_slidegen_workflow_stream(request):
                yield event
                if isinstance(event, WorkflowCompletedEvent) and event.content:
                    final_content = event.content

            if not final_content:
                yield WorkflowErrorEvent(
                    error="No content generated from workflow",
                    message="Content generation failed",
                )
                return

            markdown_doc = MarkdownDocument(final_content)

            yield StepStartedEvent(
                step_name="Presentation Export",
                message="Converting content to presentation format...",
            )

            yield ProgressEvent(
                stage="presentation_export",
                progress=0.0,
                message="Loading template...",
            )

            loop = asyncio.get_event_loop()
            template_prs = await loop.run_in_executor(None, Presentation, template)

            # Apply theme if provided
            auto_theme_llm = await get_llm_instance(request) if is_auto_theme_preset(request.theme_preset) else None
            theme = await self._resolve_theme(
                request.theme,
                request.theme_preset,
                auto_content=self._request_theme_context(request, final_content),
                auto_theme_llm=auto_theme_llm,
            )
            if theme:
                yield ProgressEvent(
                    stage="presentation_export",
                    progress=10.0,
                    message=f"Applying theme: {theme.name}...",
                )
                self.apply_theme_colors(template_prs, theme)

            yield ProgressEvent(
                stage="presentation_export",
                progress=30.0,
                message="Converting markdown to slides...",
            )

            presentation = await self.converter.generate(template_prs, markdown_doc, theme=theme)

            yield ProgressEvent(
                stage="presentation_export",
                progress=80.0,
                message="Saving presentation file...",
            )

            # Save PPTX to temp, then conditionally convert
            pptx_tmp = output_path + ".pptx"
            await loop.run_in_executor(None, partial(presentation.save, pptx_tmp))

            if request.export_as == ExportFormat.PDF:
                yield ProgressEvent(
                    stage="presentation_export",
                    progress=90.0,
                    message="Converting to PDF...",
                )
                await loop.run_in_executor(None, partial(pdf_exporter.convert, pptx_tmp, output_path))
                Path(pptx_tmp).unlink(missing_ok=True)
            else:
                Path(pptx_tmp).rename(output_path)

            yield ProgressEvent(
                stage="presentation_export",
                progress=100.0,
                message="Presentation saved successfully",
            )

            yield StepCompletedEvent(
                step_name="Presentation Export",
                content=output_path,
                message="Presentation file generated successfully",
            )

            logger.info(f"Successfully generated presentation: {output_path}")

        except Exception as e:
            logger.exception(f"Failed to generate presentation: {e}")
            yield WorkflowErrorEvent(
                error=str(e),
                message="Presentation generation failed",
            )


presentation_generator = PresentationGenerator()
