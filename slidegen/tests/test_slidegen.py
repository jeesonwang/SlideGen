import os
import sys

import pytest

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.append(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "slidegen"))

from pptx import Presentation

from slidegen.core.docparse import MarkdownDocument
from slidegen.core.docparse.markdown_parser.elements import Heading
from slidegen.core.pptgen.pptgen import PPTGen
from slidegen.core.pptgen.pptpages import CatalogPage, ChapterContentPage, ChapterHomePage, CoverPage


class TestSlideGen:
    """SlideGen核心功能测试类"""

    @pytest.fixture
    def markdown_path(self):
        """返回测试Markdown文件路径"""
        return os.path.join(os.path.dirname(__file__), "data", "report.md")

    @pytest.fixture
    def template_path(self):
        """返回测试PPT模板路径"""
        return os.path.join(os.path.dirname(__file__), "data", "DeepSeek对中国AI产业的影响.pptx")

    @pytest.fixture
    def presentation(self, template_path):
        """创建一个Presentation对象"""
        return Presentation(template_path)

    def test_markdown_document_parse(self, markdown_path):
        """测试Markdown文档解析功能"""
        # 创建MarkdownDocument对象
        doc = MarkdownDocument(markdown_path)

        # 验证文档是否正确解析
        assert doc is not None
        # 验证文档标题是否正确
        headings = [elem for elem in doc.main.children]
        assert len(headings) > 0

    @pytest.fixture
    def markdown_document(self, markdown_path):
        doc = MarkdownDocument(markdown_path)
        return doc

    @pytest.fixture
    def heading_list(self, markdown_document):
        return [elem for elem in markdown_document.main.children]

    def test_cover_page_generation(self, presentation):
        """测试PPT首页生成功能"""
        title = Heading(level=1, text="Hello World!")

        CoverPage.generate_slide(presentation, title, cover_page_index=0)

        temp_output = os.path.join(os.path.dirname(__file__), "test_cover.pptx")
        presentation.save(temp_output)

        assert os.path.exists(temp_output)

    def test_catalog_page_generation(self, presentation, heading_list):
        """测试PPT目录页生成功能"""

        CatalogPage.generate_slide(presentation, heading_list, catalog_page_index=1)
        # 保存生成的PPT文件用于测试
        temp_output = os.path.join(os.path.dirname(__file__), "test_catalog.pptx")
        presentation.save(temp_output)

        # 验证文件是否成功保存
        assert os.path.exists(temp_output)

    def test_chapter_home_page_generation(self, presentation, heading_list):
        """测试PPT章节首页生成功能"""
        title = Heading(level=2, text="Hello World! This is a test title.")
        ChapterHomePage.generate_slide(presentation, title, chapter_home_page_index=2, chapter_number=1, slide_index=2)
        temp_output = os.path.join(os.path.dirname(__file__), "test_chapter_home.pptx")
        presentation.save(temp_output)
        assert os.path.exists(temp_output)

    def test_chapter_content_page_generation(self, presentation, heading_list):
        """测试PPT章节内容页生成功能"""
        # 准备章节内容
        content = heading_list[0]
        ChapterContentPage.generate_slide(presentation, content, chapter_slide_index=4, slide_index=4)
        temp_output = os.path.join(os.path.dirname(__file__), "test_chapter_content.pptx")
        presentation.save(temp_output)
        assert os.path.exists(temp_output)

    def test_ppt_generation(self, presentation, markdown_document):
        """测试PPT生成功能"""
        ppt_gen = PPTGen()
        template_prs = Presentation(os.path.join(os.path.dirname(__file__), "data", "template_0.pptx"))
        ppt_gen.generate(template_prs, markdown_document)
        temp_output = os.path.join(os.path.dirname(__file__), "test_ppt.pptx")
        template_prs.save(temp_output)
        assert os.path.exists(temp_output)
